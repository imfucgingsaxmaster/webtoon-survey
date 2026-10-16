"""Public title-page metrics, with provenance and resumable, paced collection.

Webtoon interest counts belong to Webtoon IDs. Series downloads/comments belong
to Series product IDs. Linking retains each platform's identity. No comment text,
episode content, login cookies or private favorite state is retained.
"""
from __future__ import annotations
import argparse
import gzip
import io
import json
import re
import sqlite3
import urllib.parse
from datetime import datetime, timezone, timedelta
from decimal import Decimal
from pathlib import Path
from bs4 import BeautifulSoup
from radar import ROOT, BASE, UA, PublicClient, PublicRedirect, flattened, utcnow, public_cover_url

SERIES = 'https://series.naver.com'
METRICS = {'interest':'관심', 'downloads':'다운로드', 'comments':'댓글'}
THRESHOLDS = {'interest':[(10_000,'1만'),(100_000,'10만'),(500_000,'50만'),(1_000_000,'100만')],
              'downloads':[(100_000,'10만'),(1_000_000,'100만'),(10_000_000,'1천만'),(100_000_000,'1억')],
              'comments':[(100,'100'),(1_000,'1천'),(10_000,'1만'),(100_000,'10만')]}

def parse_count(text):
    """Preserve displayed precision; compact units are estimates, never exact."""
    raw = text.strip() if isinstance(text,str) else ''
    if not raw or raw in {'-', '—', '미표시', '관심'}:
        return {'value':None, 'display':raw, 'approximate':False}
    compact = re.sub(r'\s+', '', raw).replace(',', '')
    match = re.fullmatch(r'(?:(\d+(?:\.\d+)?)억)?(?:(\d+(?:\.\d+)?)만)?(?:(\d+(?:\.\d+)?)천)?(\d+)?(\+)?', compact)
    if not match or not any(match.groups()[:4]):
        raise ValueError(f'Unrecognized public count: {raw[:40]}')
    billion, ten_thousand, thousand, ones, plus = match.groups()
    value = sum(Decimal(n or '0')*unit for n,unit in zip((billion,ten_thousand,thousand,ones),(100000000,10000,1000,1)))
    if value != value.to_integral_value():
        raise ValueError('Non-integral count')
    return {'value':int(value), 'display':raw, 'approximate':bool(billion or ten_thousand or thousand or plus)}

def empty_metrics():
    return {k:{'value':None,'display':'','approximate':False} for k in METRICS}

def parse_webtoon(data, title_id, observed_at=None):
    if str(data.get('titleId')) != str(title_id) or not data.get('titleName'):
        raise ValueError('Webtoon title identity mismatch')
    count = data.get('favoriteCount')
    if count is not None and (type(count) is not int or count < 0):
        raise ValueError('Invalid public favoriteCount')
    contents_no=data.get('contentsNo')
    if contents_no is not None and (type(contents_no) is not int or contents_no < 0):
        raise ValueError('Invalid public contentsNo')
    tags = data.get('curationTagList')
    if tags is None: tags=[]
    if not isinstance(tags,list) or any(not isinstance(t,dict) or not isinstance(t.get('tagName'),str) for t in tags):
        raise ValueError('Unexpected public tags schema')
    metrics=empty_metrics()
    metrics['interest']={'value':count,'display':f'{count:,}' if count is not None else '', 'approximate':False}
    return {'platform':'네이버웹툰','title_id':str(title_id),'title':data['titleName'],
            'contents_no':str(contents_no) if contents_no else None,
            'cover_url':public_cover_url(data.get('thumbnailUrl')),
            'genres':sorted({t['tagName'] for t in tags if str(t.get('curationType','')).startswith('GENRE_')}),
            'official_tags':sorted({t['tagName'] for t in tags}), 'metrics':metrics,
            'source_url':f'{BASE}/webtoon/list?titleId={title_id}',
            'metadata_url':f'{BASE}/api/article/list/info?titleId={title_id}',
            'observed_at':observed_at or utcnow()}

def parse_series(html, title_id, observed_at=None):
    soup=BeautifulSoup(html,'html.parser')
    heading=soup.select_one('#content .end_head h2')
    share=soup.select_one('#spiButton2[data-url]')
    if not heading or not share or series_product_id(share['data-url']) != str(title_id):
        raise ValueError('Series product identity or public detail heading missing')
    metrics=empty_metrics()
    for field,selector in [('downloads','.end_head .btn_download span'),('comments','.end_head #commentCount')]:
        element=soup.select_one(selector)
        metrics[field]=parse_count(element.get_text(strip=True) if element else '')
    genres=sorted({a.get_text(strip=True) for a in soup.select('.end_info .info_lst a[href]')
                   if urllib.parse.urlsplit(a['href']).path=='/comic/categoryProductList.series'
                   and 'genreCode' in urllib.parse.parse_qs(urllib.parse.urlsplit(a['href']).query)})
    authors=[]
    status='미표시'
    for li in soup.select('.end_info .info_lst > ul > li'):
        text=li.get_text(' ',strip=True)
        if text.startswith(('글 ','그림 ')): authors.extend(a.get_text(strip=True) for a in li.select('a'))
        if text in ('연재중','완결','휴재'): status={'연재중':'연재'}.get(text,text)
    score=soup.select_one('.end_head .score_area em')
    rating=float(score.get_text(strip=True)) if score else None
    if rating is not None and not 0 <= rating <= 10: raise ValueError('Invalid public rating')
    cover=soup.select_one('meta[property="og:image"]')
    return {'platform':'네이버시리즈','title_id':str(title_id),'title':heading.get_text(' ',strip=True),
            'cover_url':public_cover_url(cover.get('content')) if cover else '',
            'author':' / '.join(dict.fromkeys(authors)),'status':status,'rating':rating,
            'genres':genres,'official_tags':[], 'metrics':metrics,
            'source_url':f'{SERIES}/comic/detail.series?productNo={title_id}',
            'observed_at':observed_at or utcnow()}

class SeriesLinkRestricted(PublicRedirect):
    """The official comic link is known, but its detail page requires access."""
    def __init__(self, product_id, resolution_url):
        super().__init__('Official Series link verified; public metrics restricted')
        self.product_id=str(product_id)
        self.resolution_url=resolution_url

class SeriesClient(PublicClient):
    base=SERIES
    landing='/comic/home.series'
    allowed={'/comic/home.series','/comic/top100List.series','/comic/detail.series'}

    def html(self,path,**params):
        if path not in self.allowed: raise PermissionError('Series endpoint is not allowlisted')
        url=self.base+path+('?' + urllib.parse.urlencode(params) if params else '')
        if not self.robots.can_fetch(UA,url): raise PermissionError('Series path disallowed by robots.txt')
        raw=self._read(url)
        if len(raw)>12_000_000: raise ValueError('Unexpectedly large response')
        return raw.decode('utf-8')

    def resolve_original(self,contents_no):
        """Resolve one public content ID, allowing only one canonical comic redirect."""
        if not str(contents_no).isdecimal() or int(contents_no)<1:
            raise ValueError('Invalid public content identity')
        resolution_url=f'{SERIES}/comic/detail.series?originalProductId={contents_no}'
        product_id=None
        try:
            html=self.html('/comic/detail.series',originalProductId=contents_no)
        except PublicRedirect as exc:
            # Never follow login, age, novel, arbitrary host, or redirect chains.
            product_id=series_product_id(urllib.parse.urljoin(resolution_url,exc.target)) if exc.target else None
            if not product_id:raise
            try:html=self.html('/comic/detail.series',productNo=product_id)
            except PublicRedirect:
                # Keep the already verified canonical link, never follow login.
                raise SeriesLinkRestricted(product_id,resolution_url) from None
        soup=BeautifulSoup(html,'html.parser')
        share=soup.select_one('#spiButton2[data-url]')
        ident=series_product_id(share['data-url']) if share else None
        if not ident or (product_id and ident!=product_id):
            raise ValueError('Resolved Series identity mismatch')
        row=parse_series(html,ident)
        row.update(original_product_id=str(contents_no),resolution_url=resolution_url,resolved_at=utcnow())
        return row

def series_product_id(url):
    """Strict public canonical comic identity; not a general redirect permission."""
    parsed=urllib.parse.urlsplit(url)
    query=urllib.parse.parse_qs(parsed.query,keep_blank_values=True)
    if (parsed.scheme!='https' or parsed.netloc!='series.naver.com' or
        parsed.path!='/comic/detail.series' or parsed.fragment or set(query)!={'productNo'}):return None
    values=query['productNo']
    return values[0] if len(values)==1 and values[0].isdecimal() and int(values[0])>0 else None

def series_discovery(client):
    """Follow only links shown on allowed home / daily TOP100 pages."""
    identifiers=set()
    pages=[]
    def inspect(html):
        soup=BeautifulSoup(html,'html.parser')
        for a in soup.select('#content a[href]'):
            url=urllib.parse.urlsplit(urllib.parse.urljoin(SERIES,a['href']))
            params=urllib.parse.parse_qs(url.query)
            if url.netloc!='series.naver.com': continue
            if url.path=='/comic/detail.series':
                ident=params.get('productNo',[''])[0]
                if ident.isdecimal(): identifiers.add(ident)
            elif url.path=='/comic/top100List.series' and params.get('rankingTypeCode')==['DAILY'] and params.get('categoryCode')==['ALL']:
                page=params.get('page',['1'])[0]
                if page.isdecimal() and 2<=int(page)<=5: pages.append(int(page))
    inspect(client.html('/comic/home.series'))
    inspect(client.html('/comic/top100List.series',rankingTypeCode='DAILY',categoryCode='ALL'))
    for page in sorted(set(pages)):
        inspect(client.html('/comic/top100List.series',rankingTypeCode='DAILY',categoryCode='ALL',page=page))
    if not identifiers: raise ValueError('Empty public Series discovery')
    return identifiers

def load_metadata(path=ROOT/'metadata_latest.json'):
    if not Path(path).exists(): return {'schema_version':1,'webtoon':{},'series':{},'runs':{}}
    value=json.loads(Path(path).read_text(encoding='utf-8'))
    if value.get('schema_version')!=1: raise ValueError('Unknown metadata schema')
    return value

def save_metadata(value,path):
    value['updated_at']=utcnow()
    temp=Path(path).with_suffix('.tmp')
    temp.write_text(json.dumps(value,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
    temp.replace(path)

def collect_metadata(root=ROOT,webtoon_limit=None,series_limit=None,fresh_hours=144):
    root=Path(root); path=root/'metadata_latest.json'; cache=load_metadata(path)
    unavailable=cache.setdefault('unavailable',{'webtoon':{},'series':{}})
    snapshot=json.loads((root/'catalog_latest.json').read_text(encoding='utf-8'))
    failures=[]
    cutoff=datetime.now(timezone.utc)-timedelta(hours=fresh_hours)
    def needs_refresh(record):
        return not record or datetime.fromisoformat(record['observed_at']) < cutoff
    for platform,limit in [('webtoon',webtoon_limit),('series',series_limit)]:
        if limit==0: continue
        run={'started_at':utcnow(),'state':'running','fetched':0,'requests':0,'errors':[]}
        cache['runs'][platform]=run
        save_metadata(cache,path)
        try:
            if platform=='webtoon':
                client=PublicClient()
                rank={}
                for m in snapshot['memberships']: rank[m['title_id']]=min(rank.get(m['title_id'],10**9),m['position'])
                ids={t['title_id'] for t in snapshot['titles']}
                ordered=sorted(ids,key=lambda x:(bool(cache[platform].get(x)),cache[platform].get(x,{}).get('observed_at',''),rank.get(x,10**9),int(x)))
            else:
                client=SeriesClient()
                ids=series_discovery(client) | set(cache['series'])
                ordered=sorted(ids,key=lambda x:(bool(cache[platform].get(x)),cache[platform].get(x,{}).get('observed_at',''),int(x)))
            run['target']=len(ids)
            run['scope']='정식 공개 카탈로그 전체' if platform=='webtoon' else '공개 만화 홈·일간 TOP100 및 이전 발견 작품 (시리즈 전체 아님)'
            for ident in ordered:
                if not needs_refresh(cache[platform].get(ident)) or not needs_refresh(unavailable[platform].get(ident)): continue
                if limit is not None and run['requests']>=limit: break
                run['requests']+=1
                try:
                    if platform=='webtoon': record=parse_webtoon(client.get('/api/article/list/info',titleId=ident),ident)
                    else:
                        record=parse_series(client.html('/comic/detail.series',productNo=ident),ident)
                        for key in ('original_product_id','resolution_url','resolved_at'):
                            if key in cache[platform].get(ident,{}):record[key]=cache[platform][ident][key]
                except PublicRedirect:
                    unavailable[platform][ident]={'observed_at':utcnow(),'reason':'공개 작품 정보 외 경로로 이동 요청: 로그인·연령 확인 등을 진행하지 않고 제외'}
                    continue
                unavailable[platform].pop(ident,None)
                cache[platform][ident]=record
                run['fetched']+=1
                if run['fetched']%20==0:
                    save_metadata(cache,path)
                    print(f"{platform}: {run['fetched']} refreshed; {len(cache[platform])}/{len(ids)} stored",flush=True)
            run['covered']=sum(i in cache[platform] for i in ids)
            run['fresh']=sum(not needs_refresh(cache[platform].get(i)) for i in ids)
            run['unavailable']=sum(not needs_refresh(unavailable[platform].get(i)) for i in ids)
            checked=sum(not needs_refresh(cache[platform].get(i)) or not needs_refresh(unavailable[platform].get(i)) for i in ids)
            run['state']='complete' if checked==len(ids) else 'partial'
        except Exception as exc:
            # Do not retry or bypass rate limits, authentication, robots or schema changes.
            run['state']='stopped';run['errors']=[f'{type(exc).__name__}: {str(exc)[:200]}']
            failures.append(platform)
        finally:
            run['ended_at']=utcnow();save_metadata(cache,path)
    archive=root/'metadata_history';archive.mkdir(exist_ok=True)
    stamp=cache.get('updated_at',utcnow()).replace(':','').replace('+0000','Z')
    (archive/(stamp+'.json.gz')).write_bytes(gzip.compress(json.dumps(cache,ensure_ascii=False,separators=(',',':')).encode(),mtime=0))
    export_metadata(cache,root/'outputs'/'catalogue.sqlite')
    print(json.dumps(cache['runs'],ensure_ascii=False,indent=2),flush=True)
    if failures: raise RuntimeError('Metadata collection stopped: '+', '.join(failures))
    return cache

def search_tags(record):
    tags={f'장르:{g}' for g in record.get('genres',[])}
    tags.update(f'태그:{tag}' for tag in record.get('official_tags',[]))
    for field,levels in THRESHOLDS.items():
        metric=record.get('metrics',{}).get(field,{})
        value=metric.get('value')
        if value is not None:
            for threshold,label in levels:
                if value>=threshold: tags.add(f'{METRICS[field]}:{label} 이상')
    return sorted(tags)

def build_rows(snapshot,metadata):
    rows=[]
    for item in flattened(snapshot):
        detail=metadata.get('webtoon',{}).get(item['title_id'],{})
        rows.append(dict(item,**{k:detail.get(k,default) for k,default in [('contents_no',None),('genres',[]),('official_tags',[]),('metrics',empty_metrics()),('observed_at',''),('source_url',item['url'])]}))
    for item in metadata.get('series',{}).values():
        rows.append(dict(item,url=item['source_url'],groups='공개 홈·일간 TOP100·이전 발견'))
    for row in rows:
        row['tags']=search_tags(row)
        for key in METRICS:
            metric=row['metrics'].get(key,{})
            row[key]=metric.get('value')
            row[key+'_display']=('약 ' if metric.get('approximate') else '')+(metric.get('display') or '미확인')
    return rows

def filter_rows(rows,query='',tags=(),mode='모두 포함',minimums=None):
    wanted=set(tags); result=[]; query=query.strip().casefold()
    for row in rows:
        available=set(row['tags'])
        if wanted and not (wanted<=available if mode=='모두 포함' else bool(wanted & available)): continue
        searchable=' '.join([row['title'],row.get('author',''),*row['tags']]).casefold()
        if query and query not in searchable: continue
        if any(row.get(k) is None or row[k]<v for k,v in (minimums or {}).items() if v>0): continue
        result.append(row)
    return result

def export_metadata(cache,db):
    Path(db).parent.mkdir(exist_ok=True,parents=True)
    con=sqlite3.connect(db)
    try:
        with con:
            con.execute('CREATE TABLE IF NOT EXISTS title_metrics(platform TEXT,title_id TEXT,observed_at TEXT,interest INTEGER,downloads INTEGER,comments INTEGER,genres TEXT,tags TEXT,source_url TEXT,payload TEXT, PRIMARY KEY(platform,title_id,observed_at))')
            for platform in ('webtoon','series'):
                for row in cache.get(platform,{}).values():
                    con.execute('INSERT OR IGNORE INTO title_metrics VALUES(?,?,?,?,?,?,?,?,?,?)',
                      (row['platform'],row['title_id'],row['observed_at'],*[row['metrics'][key]['value'] for key in METRICS],
                       json.dumps(row['genres'],ensure_ascii=False),json.dumps(search_tags(row),ensure_ascii=False),row['source_url'],json.dumps(row,ensure_ascii=False)))
    finally: con.close()

def export_search_excel(rows,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb=Workbook();ws=wb.active;ws.title='작품 태그 검색'
    keys=['platform','title_id','title','author','genres','interest','interest_display','downloads','downloads_display','comments','comments_display','tags','observed_at','source_url']
    ws.append(['플랫폼','작품 ID','작품명','작가','장르','관심 수','관심 원문','다운로드 수(표시값 환산)','다운로드 원문','댓글 수(표시값 환산)','댓글 원문','검색 태그','상세 관측 UTC','출처'])
    for row in rows:
        ws.append([', '.join(row.get(k,[])) if isinstance(row.get(k),list) else row.get(k) for k in keys])
    for cells in ws:
        for cell in cells:
            if isinstance(cell.value,str) and cell.value[:1] in '=+-@':cell.data_type='s'
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='123C4A')
    ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=26
    note=wb.create_sheet('지표 안내')
    for line in ['관심 수: 네이버웹툰 작품 페이지 favoriteCount 공개 집계.',
                 '다운로드·댓글 수: 네이버시리즈 작품 페이지 표시값. 만·억 등 축약 표시는 근삿값.',
                 '빈 숫자 셀은 미확인·미제공이며 0이 아님. 0은 실제 0 표시.',
                 '플랫폼과 작품 ID로 구분. 같은 제목도 자동으로 합치지 않음.',
                 '지표 태그·최소값 필터는 표시값 환산 기준. 정확한 경계 판정에는 근삿값 주의.',
                 '전체 네이버웹툰 목록을 유지. 상세 수집은 관측 시각이 있는 작품만 확인됨.',
                 '시리즈는 허용된 홈·일간 TOP100·이전 발견 작품만 포함. 전수 목록 아님.']:
        note.append([line])
    note.column_dimensions['A'].width=110
    wb.save(path)

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--webtoon-limit',type=int,default=None,help='Optional per-run detail request budget; 0 skips platform')
    parser.add_argument('--series-limit',type=int,default=None)
    parser.add_argument('--fresh-hours',type=int,default=144)
    args=parser.parse_args()
    if any(x is not None and x<0 for x in (args.webtoon_limit,args.series_limit,args.fresh_hours)): parser.error('Budgets and freshness must be nonnegative')
    collect_metadata(webtoon_limit=args.webtoon_limit,series_limit=args.series_limit,fresh_hours=args.fresh_hours)
