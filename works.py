"""Work-level search with conservative title/creator linking and source provenance."""
from __future__ import annotations
import copy
import re
import unicodedata
from collections import defaultdict
from metadata import build_rows, empty_metrics, search_tags

FORMAT_LABELS={'독점','단행본','컬러연재','연재','완결','네이버웹툰','웹툰','드라마원작','영화원작'}

def normalized(text):
    return ''.join(c for c in unicodedata.normalize('NFKC',str(text)).casefold() if c.isalnum())

def search_text(text):
    # Keep punctuation literal, while ignoring spacing and Unicode width.
    return re.sub(r'\s+','',unicodedata.normalize('NFKC',str(text)).casefold())

def base_title(title):
    def keep(match):
        return '' if match.group(1).strip() in FORMAT_LABELS else match.group(0)
    return re.sub(r'\s+',' ',re.sub(r'\[([^\]]+)\]',keep,title)).strip()

def creators(author):
    # A Han-character reading alongside a Korean pen name is an annotation.
    author=re.sub(r'(?<=[가-힣])\s*\([\u3400-\u9fff]+\)','',author or '')
    return {normalized(a) for a in re.split(r'\s*/\s*|\s*,\s*',author) if normalized(a)}

def same_creators(a,b):
    left,right=creators(a.get('author','')),creators(b.get('author',''))
    # An original novelist shared by different adaptations is insufficient.
    return bool(left and right) and (left==right or len(left & right)>=2)

def is_collected_edition(row):
    return '[단행본]' in row['title'] or bool(re.search(r'\[\d+권 세트\]',row['title']))

def build_works(snapshot,metadata):
    day_order=['월','화','수','목','금','토','일','매일+']
    catalogue_days=defaultdict(set)
    for membership in snapshot.get('memberships',[]):
        if membership['group'] in day_order:catalogue_days[membership['title_id']].add(membership['group'])
    source_rows=build_rows(snapshot,metadata)
    webtoons=[r for r in source_rows if r['platform']=='네이버웹툰']
    series=[r for r in source_rows if r['platform']=='네이버시리즈']
    groups={f"webtoon:{r['title_id']}":[r] for r in webtoons}
    # A canonical official redirect proves the venue even when its metrics
    # require login. Do not invent a Series title, author, or observation.
    restricted_links={}
    for row in webtoons:
        check=metadata.get('series_checks',{}).get(row['title_id'],{})
        product=str(check.get('series_id',''))
        contents=row.get('contents_no')
        if (check.get('state')=='linked_restricted' and contents and product.isdecimal()
            and int(product)>0 and check.get('contents_no')==contents
            and check.get('resolution_url')==f'https://series.naver.com/comic/detail.series?originalProductId={contents}'):
            restricted_links[row['title_id']]=check
    index=defaultdict(list)
    for row in webtoons:index[normalized(base_title(row['title']))].append(row)
    unmatched=[]
    for row in series:
        by_title=index[normalized(base_title(row['title']))]
        linked=[w for w in by_title if w.get('contents_no') and
                w['contents_no']==row.get('original_product_id') and
                row.get('resolution_url')==f"https://series.naver.com/comic/detail.series?originalProductId={w['contents_no']}"]
        # Individually verified title variants still require the official
        # shared identity and matching creators. Never strip season names.
        verified_alias=[]
        if row.get('verified_webtoon_title'):
            verified_alias=[w for w in index[normalized(base_title(row['verified_webtoon_title']))]
                if w.get('contents_no') and w['contents_no']==row.get('original_product_id')
                and row.get('resolution_url')==f"https://series.naver.com/comic/detail.series?originalProductId={w['contents_no']}"
                and same_creators(w,row)]
        linked=linked or verified_alias
        candidates=linked or [w for w in by_title if same_creators(w,row)]
        row['link_basis']=('공개 공통 작품번호·작가 확인' if verified_alias else '공개 공통 작품번호·제목 확인') if linked else '제목 정규화·작가 일치'
        if len(candidates)==1:groups[f"webtoon:{candidates[0]['title_id']}"].append(row)
        else:
            row['link_review']='동명·작가 일치 후보가 여러 개여서 별도 표시' if len(candidates)>1 else ''
            unmatched.append(row)
    # Series editions may share one work only when title and all creators agree.
    editions={}
    for row in unmatched:
        key=(normalized(base_title(row['title'])),tuple(sorted(creators(row.get('author','')))))
        if not key[1] or row.get('link_review'):key=(*key,row['title_id'])
        work_id=editions.setdefault(key,f"series:{row['title_id']}")
        groups.setdefault(work_id,[]).append(row)
    works=[]
    for work_id,sources in groups.items():
        sources=sorted(sources,key=lambda r:(r['platform']!='네이버웹툰',is_collected_edition(r),int(r['title_id'])))
        webtoon=next((r for r in sources if r['platform']=='네이버웹툰'),None)
        serial=next((r for r in sources if r['platform']=='네이버시리즈'),None)
        primary=webtoon or serial
        restricted=restricted_links.get(webtoon['title_id']) if webtoon and not serial else None
        metrics=empty_metrics();provenance={}
        for key,source in [('interest',webtoon),('downloads',serial),('comments',serial)]:
            if source:
                metrics[key]=copy.deepcopy(source['metrics'][key])
                provenance[key]={'platform':source['platform'],'title_id':source['title_id'],
                                 'url':source['source_url'],'observed_at':source.get('observed_at','')}
        venues=list(dict.fromkeys(r['platform'] for r in sources))
        if restricted:venues.append('네이버시리즈')
        source_states=[]
        for source in sources:
            key='webtoon' if source['platform']=='네이버웹툰' else 'series'
            blocked=metadata.get('unavailable',{}).get(key,{}).get(source['title_id'])
            if blocked and blocked['observed_at']>source.get('observed_at',''):
                has_observation=source.get('observed_at') and any(m.get('value') is not None for m in source['metrics'].values())
                source['collection_state']='기존 관측값' if has_observation else '공개 접근 제한'
            elif source.get('observed_at'):source['collection_state']='확인 완료'
            else:source['collection_state']='상세 수집 중'
            source_states.append(source['collection_state'])
        work={'work_id':work_id,'title':primary['title'] if webtoon else base_title(primary['title']),
              'publisher':serial.get('publisher','') if serial else '',
              'publisher_label':serial.get('publisher_label','') if serial else '',
              'publisher_url':serial['source_url'] if serial and serial.get('publisher') else '',
              'serial_days':(webtoon.get('serial_days') or sorted(catalogue_days[webtoon['title_id']],key=day_order.index)) if webtoon else [],
              'author':primary.get('author',''),'status':primary.get('status','미표시'),
              'rating':primary.get('rating'),'venues':venues,'sources':sources,
              'aliases':list(dict.fromkeys(r['title'] for r in sources)),
              'genres':sorted({g for r in sources for g in r.get('genres',[])}),
              'official_tags':sorted({t for r in sources for t in r.get('official_tags',[])}),
              'metrics':metrics,'metric_sources':provenance,
              'webtoon_id':webtoon['title_id'] if webtoon else '',
              'cover_url':primary.get('cover_url','') or next((r.get('cover_url') for r in sources if r.get('cover_url')),''),
              'is_new':bool(webtoon and webtoon.get('is_new')),
              'new_basis':webtoon.get('new_basis','') if webtoon else '',
              'views':webtoon.get('public_views') if webtoon else None,
              'webtoon_url':webtoon['source_url'] if webtoon else '',
              'series_url':serial['source_url'] if serial else (f"https://series.naver.com/comic/detail.series?productNo={restricted['series_id']}" if restricted else ''),
              'series_link_state':'linked_restricted' if restricted else ('linked' if serial else 'unchecked'),
              'series_link_check':restricted or {},
              'link_basis':' / '.join(dict.fromkeys(r.get('link_basis','제목 정규화·작가 일치') for r in sources if r['platform']=='네이버시리즈')) if len(sources)>1 else '단일 연재처',
              'collection_state':next((s for s in ('공개 접근 제한','상세 수집 중','기존 관측값') if s in source_states),'확인 완료')}
        work['tags']=search_tags(work)
        if restricted:work['link_basis']='공개 공통 작품번호의 공식 시리즈 주소 확인 · 지표 공개 접근 제한'
        if work['is_new']:work['tags'].insert(0,'신작')
        work['views_display']=f"{work['views']:,}" if work['views'] is not None else '미제공'
        for key,metric in metrics.items():
            work[key]=metric['value']
            if metric['value'] is not None:
                label=('약 ' if metric.get('approximate') else '')+metric['display']
            elif restricted and key in ('downloads','comments'):label='공개 접근 제한'
            elif key not in provenance:label='웹툰 연결 미확인' if key=='interest' else '시리즈 연결 미확인'
            else:
                src=webtoon if key=='interest' else serial
                label='미제공' if src.get('observed_at') else src['collection_state']
            work[key+'_display']=label
        works.append(work)
    return sorted(works,key=lambda r:(normalized(r['title']),r['work_id']))

def filter_works(works,query='',tags=(),mode='모두 포함',minimums=None,venues=()):
    wanted=set(tags);query=search_text(query);selected=set(venues)
    result=[]
    for work in works:
        if selected and not selected & set(work['venues']):continue
        if wanted and not (wanted<=set(work['tags']) if mode=='모두 포함' else bool(wanted & set(work['tags']))):continue
        fields=[work['title'],work.get('author',''),work.get('publisher',''),*work['aliases'],*work['tags']]
        if query and not any(query in search_text(field) for field in fields):continue
        if any(work[k] is None or work[k]<v for k,v in (minimums or {}).items() if v>0):continue
        result.append(work)
    return result

def query_title_ids(snapshot,query,limit=12):
    """Only a title/creator query can trigger a bounded public detail lookup."""
    query=search_text(query)
    if not query:return []
    found=[r for r in snapshot['titles'] if search_text(r['title'])==query or
           (len(query)>=2 and any(query in search_text(r[k]) for k in ('title','author')))]
    found.sort(key=lambda r:(search_text(r['title'])!=query,len(r['title']),int(r['title_id'])))
    return [r['title_id'] for r in found[:limit]]

def export_works_excel(works,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb=Workbook();ws=wb.active;ws.title='작품 통합 검색'
    ws.append(['작품 ID','작품명','연재처','작가','장르','관심 수(웹툰)','다운로드 수(시리즈)','댓글 수(시리즈)',
               '관심 원문','다운로드 원문','댓글 원문','검색 태그','확인 상태','웹툰 링크','시리즈 링크','연결 근거','조회수 공개목록 값','신작','신작 근거','제작사(시리즈 출판사)','연재요일','상태','갱신 정책'])
    detail=wb.create_sheet('연재처별 원자료')
    detail.append(['통합 작품 ID','연재처','원본 작품 ID','작품명·판본','작가','장르','관심 수','다운로드 수','댓글 수','관측 UTC','출처','공통 작품번호','공식 연결 확인 URL'])
    for w in works:
        ws.append([w['work_id'],w['title'],' / '.join(w['venues']),w['author'],', '.join(w['genres']),
                   w['interest'],w['downloads'],w['comments'],w['interest_display'],w['downloads_display'],w['comments_display'],
                   ', '.join(w['tags']),w['collection_state'],w['webtoon_url'],w['series_url'],w['link_basis'],w.get('views'),'신작' if w.get('is_new') else '',w.get('new_basis',''),w.get('publisher',''),' / '.join(w.get('serial_days',[])),w['status'],'완결 보관 · 갱신 제외' if w['status']=='완결' else '매주 갱신'])
        for s in w['sources']:
            detail.append([w['work_id'],s['platform'],s['title_id'],s['title'],s.get('author',''),', '.join(s.get('genres',[])),
                           *[s['metrics'][k]['value'] for k in ('interest','downloads','comments')],s.get('observed_at',''),s['source_url'],s.get('contents_no') or s.get('original_product_id'),s.get('resolution_url','')])
    from publishers import publisher_rankings
    rank=wb.create_sheet('제작사별 순위')
    rank.append(['제작사(시리즈 출판사)','작품 수','연재·휴재','완결 보관','다운로드 합계','다운로드 확인 작품 수','근삿값 포함'])
    for group in publisher_rankings(works):
        rank.append([group[k] for k in ('publisher','works','ongoing','finished','downloads','downloads_known','approximate')])
    note=wb.create_sheet('지표 안내')
    for line in ['공개 공통 작품번호와 제목으로 확인하거나, 동일 제목에 작가가 모두 또는 2명 이상 일치할 때 통합합니다. 동명 후보가 여러 개면 별도로 표시합니다.',
                 '관심은 웹툰, 다운로드·댓글은 시리즈 대표 연재판의 값입니다. 서로 다른 플랫폼·판본의 수치는 합산하지 않습니다.',
                 '시리즈에 여러 판본이 있으면 연재판을 우선하며, 모든 판본의 원자료와 출처는 별도 시트에 보존합니다.',
                 '빈 숫자는 연재처 연결 미확인·미제공·수집 중·공개 접근 제한이며 0이 아닙니다. 연결 미확인은 해당 플랫폼에 미연재라는 뜻이 아닙니다.',
                 '만·억 등 축약된 원문은 근삿값입니다. 태그·최소값 검색도 표시값 환산 기준입니다.',
                 '제작사는 시리즈 페이지의 출판사 기준이며 대표 시리즈 판본을 작품당 한 번만 합산합니다. 미확인 수치는 합계에서 제외합니다.',
                 '완결 작품은 마지막 관측을 보관하며 갱신하지 않습니다. 완결 보관 수치는 제작사 합계에 포함됩니다. 제작사 미확인은 순위 밖의 별도 그룹입니다.']:
        note.append([line])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str) and cell.value[:1] in '=+-@':cell.data_type='s'
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='123C4A')
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=27
    note.column_dimensions['A'].width=120
    wb.save(path)
