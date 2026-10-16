import copy
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
import radar
import metadata
import works
import live_data
import tests as fixtures

class WorkTests(unittest.TestCase):
    def setUp(self):
        self.snapshot=radar.collect(fixtures.FakeClient())
        self.snapshot['titles'][0].update(title='같은 제목',author='글작가 / 그림작가')
        webtoon=fixtures.MetadataTests().webtoon()
        series=fixtures.MetadataTests().series()
        series.update(author='그림작가, 글작가')
        self.cache={'schema_version':1,'webtoon':{'1':webtoon},'series':{'1':series}}

    def test_matching_work_combines_venues_and_metric_provenance(self):
        result=works.build_works(self.snapshot,self.cache)
        self.assertEqual(len(result),10)
        row=next(w for w in result if w['webtoon_id']=='1')
        self.assertEqual(row['venues'],['네이버웹툰','네이버시리즈'])
        self.assertEqual((row['interest'],row['downloads'],row['comments']),(100000,286830000,14000))
        self.assertEqual(row['metric_sources']['interest']['platform'],'네이버웹툰')
        self.assertEqual(row['metric_sources']['downloads']['platform'],'네이버시리즈')
        self.assertEqual(len(works.filter_works(result,tags=['관심:10만 이상','댓글:1만 이상'])),1)
        self.assertEqual(len(works.filter_works(result,venues=['네이버시리즈'])),1)
        self.assertEqual(len(works.filter_works(result,minimums={'interest':100001})),0)

    def test_same_title_with_different_creators_is_not_merged(self):
        self.cache['series']['1']['author']='다른 작가'
        self.assertEqual(len(works.build_works(self.snapshot,self.cache)),11)
        self.cache['series']['1']['author']='글작가 / 다른그림작가'
        self.assertEqual(len(works.build_works(self.snapshot,self.cache)),11)

    def test_ambiguous_and_missing_authors_stay_separate(self):
        self.snapshot['titles'][1].update(title='같은 제목',author='글작가 / 그림작가')
        self.assertEqual(len(works.build_works(self.snapshot,self.cache)),11)
        self.cache['series']['1']['author']=''
        self.assertEqual(len(works.build_works(self.snapshot,self.cache)),11)

    def test_editions_preserve_metrics_without_summing(self):
        edition=copy.deepcopy(self.cache['series']['1']);edition.update(title_id='22',title='같은 제목 [단행본]')
        edition['metrics']['downloads']=metadata.parse_count('100')
        self.cache['series']['22']=edition
        result=works.build_works(self.snapshot,self.cache)
        row=next(w for w in result if w['webtoon_id']=='1')
        self.assertEqual(len(result),10);self.assertEqual(len(row['sources']),3)
        self.assertEqual(row['downloads'],286830000)
        self.assertEqual(len(works.filter_works(result,query='단행본')),1)
        edition['title']='같은 제목 시즌2'
        self.assertEqual(len(works.build_works(self.snapshot,self.cache)),11)

    def test_unobserved_zero_and_not_provided_are_distinct(self):
        self.cache['webtoon']['1']=fixtures.MetadataTests().webtoon(count=0)
        result=works.build_works(self.snapshot,self.cache)
        row=next(w for w in result if w['webtoon_id']=='1')
        self.assertEqual(row['interest_display'],'0')
        pending=next(w for w in result if w['webtoon_id']=='2')
        self.assertEqual(pending['interest_display'],'상세 수집 중')
        self.assertEqual(pending['downloads_display'],'시리즈 연결 미확인')
        self.cache['webtoon']['1']=fixtures.MetadataTests().webtoon(count=None)
        row=next(w for w in works.build_works(self.snapshot,self.cache) if w['webtoon_id']=='1')
        self.assertEqual(row['interest_display'],'미제공')

    def test_excel_contains_unified_rows_and_every_original_source(self):
        self.snapshot['titles'][0]['title']='=1+1'
        self.cache['series']['1']['title']='=1+1'
        rows=works.build_works(self.snapshot,self.cache);buf=io.BytesIO()
        works.export_works_excel(rows,buf)
        wb=load_workbook(io.BytesIO(buf.getvalue()))
        self.assertEqual(wb['작품 통합 검색'].max_row,11)
        self.assertEqual(wb['연재처별 원자료'].max_row,12)
        cells=[r[1] for r in wb['작품 통합 검색'].iter_rows(min_row=2) if r[1].value=='=1+1']
        self.assertEqual(len(cells),1);self.assertEqual(cells[0].data_type,'s')

    def test_query_lookup_is_bounded_and_literal(self):
        self.assertEqual(works.query_title_ids(self.snapshot,'같은 제목'),['1'])
        self.assertEqual(works.query_title_ids(self.snapshot,'['),[])
        self.assertLessEqual(len(works.query_title_ids(self.snapshot,'Test',limit=2)),2)
        self.snapshot['titles'][0]['title']='숲'
        self.assertEqual(works.query_title_ids(self.snapshot,'숲'),['1'])

    def test_han_pen_name_annotation_matches_without_fuzzy_author_matching(self):
        self.assertEqual(works.creators('운(雲) / 배민기'),works.creators('운 / 배민기'))
        self.assertNotEqual(works.creators('삼돌'),works.creators('삼둘'))

class LiveDataTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)
        self.environment=patch.dict(os.environ,{'WEBTOON_RADAR_OFFLINE':'0'});self.environment.start();self.addCleanup(self.environment.stop)
        self.snapshot=radar.collect(fixtures.FakeClient())
        live_data.write_json(self.root/'catalog_latest.json',self.snapshot)
        self.row=fixtures.MetadataTests().webtoon()

    def test_newer_records_survive_merge_and_old_exclusion_is_cleared(self):
        older=copy.deepcopy(self.row);older['observed_at']='2026-10-14T00:00:00+00:00'
        result=live_data.merge_metadata({'webtoon':{'1':self.row}},
            {'webtoon':{'1':older},'unavailable':{'webtoon':{'1':{'observed_at':older['observed_at']}}}})
        self.assertEqual(result['webtoon']['1'],self.row)
        self.assertNotIn('1',result['unavailable']['webtoon'])

    def test_search_fills_once_and_reuses_public_observation(self):
        class Client:
            calls=0
            def get(self,path,**params):
                Client.calls+=1
                return {'titleId':int(params['titleId']),'titleName':'공개 작품','favoriteCount':123,'curationTagList':[]}
        with patch.object(live_data.time,'sleep'):
            first=live_data.fill_searched_titles(['1','1'],self.root,Client)
            second=live_data.fill_searched_titles(['1'],self.root,Client)
        self.assertEqual((first['fetched'],second['fetched'],Client.calls),(1,0,1))
        self.assertEqual(live_data.load_live_metadata(self.root)['webtoon']['1']['metrics']['interest']['value'],123)

    def test_rate_failure_stops_and_cools_down_without_losing_previous(self):
        live_data.write_json(self.root/'metadata_latest.json',{'schema_version':1,'webtoon':{'1':self.row},'series':{}})
        class Client:
            calls=0
            def get(self,*a,**k):Client.calls+=1;raise PermissionError('rate limited')
        with patch.object(live_data.time,'sleep'):
            result=live_data.fill_searched_titles(['2','3'],self.root,Client)
            live_data.fill_searched_titles(['2'],self.root,Client)
        self.assertEqual(result['state'],'temporarily_unavailable');self.assertEqual(Client.calls,1)
        self.assertIn('1',live_data.load_live_metadata(self.root)['webtoon'])

    def test_login_redirect_is_recorded_without_following(self):
        class Client:
            def get(self,*a,**k):raise radar.PublicRedirect('login')
        with patch.object(live_data.time,'sleep'):
            result=live_data.fill_searched_titles(['2'],self.root,Client)
        self.assertEqual(result['fetched'],0)
        self.assertIn('2',live_data.load_live_metadata(self.root)['unavailable']['webtoon'])

    def test_public_sync_is_throttled_and_preserves_newer_local_details(self):
        live_data.write_json(self.root/'metadata_latest.json',{'schema_version':1,'webtoon':{'1':self.row},'series':{}})
        calls=[];older=copy.deepcopy(self.row);older['observed_at']='2026-10-14T00:00:00+00:00'
        def fetch(url):
            calls.append(url)
            if '/git/ref/' in url:return {'object':{'sha':'a'*40}}
            if url.endswith('catalog_latest.json'):return self.snapshot
            return {'schema_version':1,'webtoon':{'1':older},'series':{}}
        self.assertTrue(live_data.sync_public_observations(self.root,fetcher=fetch)['updated'])
        self.assertFalse(live_data.sync_public_observations(self.root,fetcher=fetch)['updated'])
        self.assertEqual(len(calls),3)
        self.assertEqual(live_data.load_live_metadata(self.root)['webtoon']['1'],self.row)
        self.assertFalse(live_data.sync_public_observations(self.root,force=True,fetcher=fetch)['updated'])
        self.assertEqual(len(calls),4)

    def test_malformed_remote_snapshot_does_not_replace_local(self):
        before=(self.root/'catalog_latest.json').read_bytes()
        def fetch(url):return {'object':{'sha':'b'*40}} if '/git/ref/' in url else {'schema_version':99}
        self.assertEqual(live_data.sync_public_observations(self.root,fetcher=fetch)['state'],'unavailable')
        self.assertEqual((self.root/'catalog_latest.json').read_bytes(),before)

if __name__=='__main__':unittest.main()
